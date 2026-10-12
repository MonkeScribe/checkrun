from datetime import date
from decimal import Decimal

import openpyxl
import pandas as pd
import pytest

from checkrun import core


def raw_file(tmp_path, rows, name="raw.xlsx"):
    p = tmp_path / name
    pd.DataFrame(rows).to_excel(p, index=False)
    return p


GOOD = [
    {"Vendor": "Acme Motors", "Check Date": "2026-08-15", "Total": "$1,234.50", "Check #": 1001, "Description": "Aug"},
    {"Vendor": "Bob's Parts", "Check Date": "2026-08-15", "Total": 99, "Check #": 1002},
]


def test_end_to_end(tmp_path):
    filled, iif = tmp_path / "f.xlsx", tmp_path / "o.iif"
    assert core.fill_template(raw_file(tmp_path, GOOD), filled) == 2
    ws = openpyxl.load_workbook(filled)["Check Data"]
    assert ws["B3"].value == "Acme Motors" and ws["C3"].value == "Checking" and ws["D4"].value == "Dealer Checks"
    assert core.convert_template_to_iif(filled, iif) == 2

    lines = iif.read_bytes().decode("cp1252").split("\r\n")
    assert lines[0].startswith("!TRNS\tTRNSID") and lines[2] == "!ENDTRNS"
    assert lines[3] == "TRNS\t\tCHECK\t08/15/2026\tChecking\tAcme Motors\t\t-1234.50\t1001\tAug\tY"
    assert lines[4] == "SPL\t\tCHECK\t08/15/2026\tDealer Checks\t\t1234.50\t1001\tAug"
    assert lines[5] == "ENDTRNS"
    assert lines[6].endswith("-99.00\t1002\t\tY")


def test_blocks_balance():
    c = core.Check(date(2026, 1, 2), "X", "Checking", "Dealer Checks", Decimal("10.10"), "1", "", "")
    t = [l.split("\t") for l in core.build_iif([c]).splitlines() if l.startswith(("TRNS", "SPL"))]
    assert Decimal(t[0][7]) + Decimal(t[1][6]) == 0


def test_missing_column(tmp_path):
    with pytest.raises(core.CheckRunError, match="missing column"):
        core.fill_template(raw_file(tmp_path, [{"Vendor": "A", "Total": 1}]), tmp_path / "f.xlsx")


def test_bad_rows_reported_together(tmp_path):
    rows = GOOD + [
        {"Vendor": "", "Check Date": "nope", "Total": -5, "Check #": 1001},
    ]
    with pytest.raises(core.CheckRunError) as e:
        core.fill_template(raw_file(tmp_path, rows), tmp_path / "f.xlsx")
    msg = str(e.value)
    assert "Row 4: invalid or missing Date" in msg and "missing Payee" in msg
    assert "Amount must be a positive" in msg and "duplicate Check Number 1001" in msg


def test_tabs_and_newlines_stripped(tmp_path):
    rows = [{"Vendor": "A\tB\nC", "Check Date": "2026-01-01", "Total": 5, "Check #": 7}]
    filled, iif = tmp_path / "f.xlsx", tmp_path / "o.iif"
    core.fill_template(raw_file(tmp_path, rows), filled)
    core.convert_template_to_iif(filled, iif)
    assert "A B C" in iif.read_text("cp1252")


def test_gap_in_filled_template_is_caught(tmp_path):
    filled = tmp_path / "f.xlsx"
    core.fill_template(raw_file(tmp_path, GOOD), filled)
    wb = openpyxl.load_workbook(filled)
    wb["Check Data"]["B3"].value = None
    wb.save(filled)
    with pytest.raises(core.CheckRunError, match="no Payee"):
        core.read_template_checks(filled)


def test_unencodable_character(tmp_path):
    rows = [{"Vendor": "日本", "Check Date": "2026-01-01", "Total": 5, "Check #": 7}]
    filled = tmp_path / "f.xlsx"
    core.fill_template(raw_file(tmp_path, rows), filled)
    with pytest.raises(core.CheckRunError, match="Windows-1252"):
        core.convert_template_to_iif(filled, tmp_path / "o.iif")


def test_csv_input(tmp_path):
    p = tmp_path / "r.csv"
    p.write_text("Payee,Date,Amount,Check Number\nZed,8/1/2026,12.5,55\n")
    assert core.fill_template(p, tmp_path / "f.xlsx") == 1
