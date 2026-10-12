"""Core logic: raw data -> filled template -> IIF.

Kept free of any GUI code so it can be tested and scripted on its own.
The IIF layout mirrors the formulas in templates/Check_to_IIF_Template.xlsx.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import pandas as pd

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "templates" / "Check_to_IIF_Template.xlsx"
DATA_SHEET = "Check Data"
FIRST_DATA_ROW = 3  # row 2 holds the default bank/expense accounts
TEMPLATE_CAPACITY = 60  # rows the template's own "IIF Output" tab can cover

# Check Data columns: A..H
FIELDS = ["date", "payee", "bank", "expense", "amount", "check_no", "memo", "class_"]
COLUMN_OF = {name: i for i, name in enumerate(FIELDS, start=1)}

# Accepted header spellings in the raw file (compared after normalising).
ALIASES = {
    "date": ["date", "check date", "payment date", "pay date"],
    "payee": ["payee", "vendor", "name", "pay to", "payto", "paid to"],
    "bank": ["bank account", "bank", "account"],
    "expense": ["expense account", "expense", "category", "gl account"],
    "amount": ["amount", "check amount", "total", "payment amount", "payment"],
    "check_no": ["check number", "check no", "check num", "check #", "check", "number", "num", "doc num", "docnum"],
    "memo": ["memo", "description", "note", "notes"],
    "class_": ["class"],
}
REQUIRED = ["date", "payee", "amount", "check_no"]

IIF_HEADER = (
    "!TRNS\tTRNSID\tTRNSTYPE\tDATE\tACCNT\tNAME\tCLASS\tAMOUNT\tDOCNUM\tMEMO\tTOPRINT\n"
    "!SPL\tSPLID\tTRNSTYPE\tDATE\tACCNT\tCLASS\tAMOUNT\tDOCNUM\tMEMO\n"
    "!ENDTRNS\n"
)


class CheckRunError(Exception):
    """User-facing problem with the input data. Message lists every issue found."""


@dataclass
class Check:
    date: date
    payee: str
    bank: str
    expense: str
    amount: Decimal
    check_no: str
    memo: str
    class_: str


def _norm(s: object) -> str:
    return re.sub(r"\s+", " ", str(s).strip().lower())


def _clean(v: object) -> str:
    """Cell value -> single-line text safe for a tab-delimited file."""
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # 1001.0 -> "1001"
    return re.sub(r"[\t\r\n]+", " ", str(v)).strip()


def _to_date(v: object) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if v is None or _clean(v) == "":
        return None
    try:
        ts = pd.to_datetime(v)
    except (ValueError, TypeError):
        return None
    return None if pd.isna(ts) else ts.date()


def _to_amount(v: object) -> Decimal | None:
    s = _clean(v).replace("$", "").replace(",", "")
    if s.startswith("(") and s.endswith(")"):
        s = "-" + s[1:-1]
    try:
        d = Decimal(s)
    except InvalidOperation:
        return None
    return d if d.is_finite() else None


# --------------------------------------------------------------------------
# Step 1: raw data -> filled template
# --------------------------------------------------------------------------
def read_raw(path: str | Path) -> pd.DataFrame:
    """Load the raw file and rename its columns to the canonical field names."""
    path = Path(path)
    if path.suffix.lower() in {".csv", ".txt"}:
        df = pd.read_csv(path, dtype=object, keep_default_na=False)
    else:
        df = pd.read_excel(path, dtype=object)
    df = df.dropna(how="all")

    rename: dict[str, str] = {}
    for col in df.columns:
        for field, names in ALIASES.items():
            if _norm(col) in names and field not in rename.values():
                rename[col] = field
                break
    df = df.rename(columns=rename)

    missing = [f for f in REQUIRED if f not in df.columns]
    if missing:
        raise CheckRunError(
            "The raw file is missing column(s): "
            + ", ".join(missing)
            + f".\nColumns found: {', '.join(map(str, df.columns))}"
        )
    return df.reset_index(drop=True)


def fill_template(raw_path: str | Path, out_path: str | Path, template_path: str | Path = TEMPLATE_PATH) -> int:
    """Copy raw rows into the template's Check Data tab and save it. Returns row count."""
    df = read_raw(raw_path)
    checks = _validate(_rows_from_raw(df, _template_defaults(template_path)))

    wb = openpyxl.load_workbook(template_path)
    ws = wb[DATA_SHEET]
    last = max(ws.max_row, FIRST_DATA_ROW + len(checks) - 1)
    for r in range(FIRST_DATA_ROW, last + 1):  # wipe old entries (keeps formatting)
        for col in (1, 2, 5, 6, 7, 8):
            ws.cell(r, col).value = None
    for i, c in enumerate(checks):
        r = FIRST_DATA_ROW + i
        ws.cell(r, COLUMN_OF["date"]).value = c.date
        ws.cell(r, COLUMN_OF["date"]).number_format = "mm/dd/yyyy"
        ws.cell(r, COLUMN_OF["payee"]).value = c.payee
        ws.cell(r, COLUMN_OF["bank"]).value = c.bank
        ws.cell(r, COLUMN_OF["expense"]).value = c.expense
        ws.cell(r, COLUMN_OF["amount"]).value = float(c.amount)
        ws.cell(r, COLUMN_OF["amount"]).number_format = "#,##0.00"
        ws.cell(r, COLUMN_OF["check_no"]).value = int(c.check_no) if c.check_no.isdigit() else c.check_no
        ws.cell(r, COLUMN_OF["memo"]).value = c.memo or None
        ws.cell(r, COLUMN_OF["class_"]).value = c.class_ or None
    wb.calculation.fullCalcOnLoad = True  # template formulas recalc when opened in Excel
    wb.save(out_path)
    return len(checks)


def _template_defaults(template_path: str | Path) -> dict[str, str]:
    ws = openpyxl.load_workbook(template_path, read_only=True)[DATA_SHEET]
    row2 = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    return {"bank": _clean(row2[2]), "expense": _clean(row2[3])}


def _rows_from_raw(df: pd.DataFrame, defaults: dict[str, str]) -> list[tuple[int, dict]]:
    out = []
    for i, rec in enumerate(df.to_dict("records")):
        row = {f: rec.get(f) for f in FIELDS}
        for f in ("bank", "expense"):
            if not _clean(row[f]):
                row[f] = defaults[f]
        out.append((i + 2, row))  # +2: header is spreadsheet row 1
    return out


# --------------------------------------------------------------------------
# Step 2: filled template -> IIF
# --------------------------------------------------------------------------
def read_template_checks(path: str | Path) -> list[Check]:
    """Read checks from a filled template's Check Data tab."""
    ws = openpyxl.load_workbook(path, data_only=True)[DATA_SHEET]
    rows, gap_row = [], None
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 9)]
        row = dict(zip(FIELDS, vals))
        if not _clean(row["payee"]):
            # defaults in cols C/D don't count as data
            if any(_clean(row[f]) for f in ("date", "amount", "check_no", "memo", "class_")):
                gap_row = gap_row or r
            continue
        if gap_row:
            raise CheckRunError(
                f"Row {gap_row} has data but no Payee, and more checks follow it. "
                "Fill the Payee or clear the row so no checks are skipped."
            )
        rows.append((r, row))
    return _validate(rows)


def _validate(rows: list[tuple[int, dict]]) -> list[Check]:
    checks: list[Check] = []
    errors: list[str] = []
    seen: dict[str, int] = {}
    for rownum, row in rows:
        if not _clean(row["payee"]) and not any(_clean(row[f]) for f in ("date", "amount", "check_no")):
            continue  # fully empty raw row
        where = f"Row {rownum}"
        d = _to_date(row["date"])
        amt = _to_amount(row["amount"])
        payee, bank, exp = _clean(row["payee"]), _clean(row["bank"]), _clean(row["expense"])
        no = _clean(row["check_no"])
        if d is None:
            errors.append(f"{where}: invalid or missing Date ({_clean(row['date'])!r})")
        if not payee:
            errors.append(f"{where}: missing Payee")
        if amt is None or amt <= 0:
            errors.append(f"{where}: Amount must be a positive number ({_clean(row['amount'])!r})")
        elif amt != amt.quantize(Decimal("0.01")):
            errors.append(f"{where}: Amount has more than 2 decimal places ({amt})")
        if not no:
            errors.append(f"{where}: missing Check Number")
        elif no in seen:
            errors.append(f"{where}: duplicate Check Number {no} (also row {seen[no]})")
        else:
            seen[no] = rownum
        if not bank or not exp:
            errors.append(f"{where}: missing Bank Account or Expense Account")
        if d and payee and amt is not None and amt > 0 and no and bank and exp:
            checks.append(Check(d, payee, bank, exp, amt, no, _clean(row["memo"]), _clean(row["class_"])))
    if errors:
        shown = errors[:25]
        more = f"\n...and {len(errors) - 25} more." if len(errors) > 25 else ""
        raise CheckRunError("Please fix these problems:\n" + "\n".join(shown) + more)
    if not checks:
        raise CheckRunError("No checks found.")
    return checks


def summarize(checks: list[Check], today: date | None = None) -> str:
    """Human-readable recap used for the 'are these dates right?' confirmation."""
    today = today or date.today()
    dates = sorted({c.date for c in checks})
    total = sum(c.amount for c in checks)
    rng = f"{dates[0]:%m/%d/%Y}" if len(dates) == 1 else f"{dates[0]:%m/%d/%Y} – {dates[-1]:%m/%d/%Y}"
    lines = [
        f"Checks: {len(checks)}",
        f"Check date(s): {rng}",
        f"Total: ${total:,.2f}",
        f"Today is {today:%m/%d/%Y}.",
    ]
    odd = [c for c in checks if abs((c.date - today).days) > 30]
    if odd:
        lines.append(f"\n⚠ {len(odd)} check(s) are dated more than 30 days from today (e.g. check #{odd[0].check_no}: {odd[0].date:%m/%d/%Y}).")
    return "\n".join(lines)


def build_iif(checks: list[Check]) -> str:
    """Return the full IIF text (LF line endings, written as CRLF on save)."""
    lines = [IIF_HEADER.rstrip("\n")]
    for c in checks:
        d = c.date.strftime("%m/%d/%Y")
        amt = f"{c.amount:.2f}"
        # SPL deliberately has no NAME: a payee there is read as Customer:Job (error 10143).
        lines.append("\t".join(["TRNS", "", "CHECK", d, c.bank, c.payee, c.class_, "-" + amt, c.check_no, c.memo, "Y"]))
        lines.append("\t".join(["SPL", "", "CHECK", d, c.expense, c.class_, amt, c.check_no, c.memo]))
        lines.append("ENDTRNS")
    return "\n".join(lines) + "\n"


def write_iif(checks: list[Check], out_path: str | Path) -> None:
    text = build_iif(checks)
    try:
        data = text.encode("cp1252")  # what QuickBooks Desktop expects
    except UnicodeEncodeError as e:
        bad = text[e.start : e.end]
        raise CheckRunError(
            f"Character {bad!r} can't be saved in the Windows-1252 encoding QuickBooks needs. "
            "Replace it in the spreadsheet."
        ) from None
    Path(out_path).write_bytes(data.replace(b"\n", b"\r\n"))


def convert_template_to_iif(template_path: str | Path, out_path: str | Path) -> int:
    checks = read_template_checks(template_path)
    write_iif(checks, out_path)
    return len(checks)
